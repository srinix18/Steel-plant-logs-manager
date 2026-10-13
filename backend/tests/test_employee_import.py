"""HR bulk upload of employees from an Excel file: template, validation, commit."""

from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.import_engine.handlers.employees import COLUMNS

HEADERS = ["employee_uid", "full_name", "email", "password", "role", "department", "process", "maintenance_division", "shift", "date_of_joining", "basic", "hra"]


def _sheet(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(HEADERS)  # "department" is an accepted heading for the department column
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _run(api, who, content):
    up = await api.post(who, "/imports/employees/upload", files={"file": ("staff.xlsx", content, "application/octet-stream")})
    assert up.status_code == 200, up.text
    job = up.json()["id"]
    await api.ok(who, "POST", f"/imports/jobs/{job}/preview")
    validated = await api.ok(who, "POST", f"/imports/jobs/{job}/validate")
    return job, validated


async def test_template_has_guide_lists_and_dropdowns(api):
    r = await api.get("hr", "/imports/employees/template")
    assert r.status_code == 200
    wb = load_workbook(BytesIO(r.content))
    assert {"Read me first", "Employees", "Lists", "Example"} <= set(wb.sheetnames)
    assert [c.value for c in wb["Employees"][1]] == COLUMNS
    codes = {row[1].value for row in wb["Lists"].iter_rows(min_row=2) if row[1].value}
    assert "SMS" in codes
    assert wb["Employees"].data_validations.dataValidation  # role / department dropdowns


async def test_import_creates_users_with_shift_and_salary(api):
    content = _sheet([
        ["IMP-001", "Import Worker", "import.worker@example.com", "Temp@1234", "worker", "SMS", "", "", "A", "2025-01-06", 21000, 6500],
        ["IMP-002", "Import Fitter", "import.fitter@example.com", "Temp@1234", "maintenance", "SMS", "", "equipment", "", "06-01-2025", "", ""],
        ["IMP-003", "Bad Role", "bad.role@example.com", "Temp@1234", "astronaut", "SMS", "", "", "", "", "", ""],
        ["IMP-004", "No Dept", "no.dept@example.com", "Temp@1234", "worker", "", "", "", "", "", "", ""],
        ["IMP-005", "Same Mail", "import.worker@example.com", "Temp@1234", "worker", "SMS", "", "", "", "", "", ""],
        ["IMP-006", "Existing", "melter@chandansteel.com", "Temp@1234", "worker", "SMS", "", "", "", "", "", ""],
        ["IMP-007", "No Process", "no.process@example.com", "Temp@1234", "supervisor", "SMS", "", "", "B", "", "", ""],
        ["IMP-008", "Iaf Supervisor", "iaf.new@example.com", "Temp@1234", "supervisor", "SMS", "IAF", "", "B", "", "", ""],
    ])
    job, validated = await _run(api, "hr", content)
    assert validated["valid_count"] == 3
    assert validated["invalid_count"] == 5
    reasons = {r["row_number"]: " ".join(r["errors"]) for r in validated["rows"] if r["errors"]}
    assert "Invalid role" in reasons[4]
    assert "department" in reasons[5]
    assert "Duplicate email" in reasons[6]
    assert "already registered" in reasons[7]
    assert "process is required" in reasons[8]

    done = await api.ok("hr", "POST", f"/imports/jobs/{job}/commit")
    jobinfo = await api.ok("hr", "GET", f"/imports/jobs/{job}")
    assert done["imported"] == 3 and done["failed"] == 0, [r["errors"] for r in jobinfo["rows"] if r["errors"]]

    # The new people can log in, and the optional columns took effect.
    login = await api.post(None, "/auth/login", json={"email": "import.worker@example.com", "password": "Temp@1234"})
    assert login.status_code == 200
    employees = await api.ok("hr", "GET", "/workforce/employees")
    worker = next(e for e in employees if e["email"] == "import.worker@example.com")
    assert worker["employee_uid"] == "IMP-001"
    assigned = await api.ok("hr", "GET", "/workforce/shift-assignments")
    assert any(a["user_id"] == worker["id"] and a["shift_code"] == "A" for a in assigned)
    structures = await api.ok("hr", "GET", "/workforce/payroll/salary-structures", params={"user_id": worker["id"]})
    assert structures and structures[0]["basic"] == 21000
    fitter = next(e for e in employees if e["email"] == "import.fitter@example.com")
    assert fitter["role"] == "maintenance"


async def test_workers_cannot_import(api):
    r = await api.post("worker_sms", "/imports/employees/upload", files={"file": ("x.xlsx", _sheet([]), "application/octet-stream")})
    assert r.status_code == 403
