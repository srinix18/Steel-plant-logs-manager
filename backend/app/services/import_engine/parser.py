from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

DATA_SHEET = "Employees"


def build_template(columns: list[str]) -> bytes:
    """Plain header-only workbook, for modules without a richer template."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Import"
    for idx, col in enumerate(columns, start=1):
        ws.cell(row=1, column=idx, value=col)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def normalize_header(raw: Any, aliases: dict[str, str] | None = None) -> str:
    key = str(raw).strip().lower().replace("*", "").replace("-", "_").replace(" ", "_") if raw is not None else ""
    key = key.strip("_")
    return (aliases or {}).get(key, key)


def _clean(val: Any) -> Any:
    """Make a cell JSON-safe and tidy: dates become ISO text, whole floats become ints, text is trimmed."""
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.date().isoformat() if val.time() == datetime.min.time() else val.isoformat()
    if isinstance(val, date):
        return val.isoformat()
    if isinstance(val, float) and val.is_integer():
        return int(val)
    if isinstance(val, str):
        return val.strip()
    return val


def parse_xlsx(content: bytes, aliases: dict[str, str] | None = None) -> list[dict[str, Any]]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb[DATA_SHEET] if DATA_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [normalize_header(h, aliases) for h in header_row]
    result: list[dict[str, Any]] = []
    for row in rows_iter:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        item: dict[str, Any] = {}
        for i, key in enumerate(headers):
            if not key:
                continue
            item[key] = _clean(row[i] if i < len(row) else None)
        result.append(item)
    return result
