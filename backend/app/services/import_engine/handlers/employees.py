from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Department, Shift, User
from app.models.enums import EmploymentStatus, EmploymentType, ObservationCategory, UserRole
from app.schemas.moi import ShiftAssignmentCreate, WorkforceEmployeeCreate
from app.schemas.workforce_ops import SalaryStructureCreate
from app.services.import_engine.registry import register_module
from app.services.payroll_service import PayrollService
from app.services.workforce_service import WorkforceService

COLUMNS = [
    "employee_uid",
    "full_name",
    "email",
    "password",
    "role",
    "department_id",
    "maintenance_division",
    "shift",
    "designation",
    "phone",
    "employment_type",
    "employment_status",
    "date_of_joining",
    "basic",
    "hra",
    "allowances",
    "pf",
    "esi",
]

REQUIRED = ["full_name", "email", "password", "role"]

ALIASES = {
    "department": "department_id",
    "department_code": "department_id",
    "dept": "department_id",
    "name": "full_name",
    "employee_name": "full_name",
    "employee_id": "employee_uid",
    "emp_id": "employee_uid",
    "employee_code": "employee_uid",
    "mobile": "phone",
    "phone_number": "phone",
    "mobile_number": "phone",
    "joining_date": "date_of_joining",
    "doj": "date_of_joining",
    "basic_salary": "basic",
    "division": "maintenance_division",
    "shift_code": "shift",
}

# Roles an HR user can load from a file; platform roles (admin, CEO) are created by hand.
IMPORTABLE_ROLES = ["worker", "supervisor", "hod", "maintenance", "hr"]
NO_DEPARTMENT_ROLES = {"hr"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

_workforce = WorkforceService()
_payroll = PayrollService()


async def _resolve_department(
    session: AsyncSession, org_id: UUID, dept_raw: Any
) -> tuple[Department | None, list[str]]:
    """Resolve the department column: a department code (SMS), a department name, or a UUID."""
    if not dept_raw:
        return None, []
    raw = str(dept_raw).strip()
    try:
        dept = await session.get(Department, UUID(raw))
        if dept and dept.organisation_id == org_id:
            return dept, []
        return None, ["Department not found"]
    except ValueError:
        pass
    result = await session.execute(select(Department).where(Department.organisation_id == org_id))
    for dept in result.scalars():
        if raw.upper() == (dept.code or "").upper() or raw.lower() == (dept.name or "").lower():
            return dept, []
    return None, [f"Department not found: {raw} (use the department code from the Lists sheet)"]


def _parse_date(val: Any) -> date | None:
    if val is None or val == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_role(val: Any) -> UserRole | None:
    s = str(val or "").strip().lower().replace(" ", "_")
    if s in ("head_of_department", "hod"):
        s = "hod"
    if s in ("shift_incharge", "shift_in_charge", "incharge"):
        s = "supervisor"
    if s in ("operator", "employee", "member"):
        s = "worker"
    return next((r for r in UserRole if r.value == s and r.value in IMPORTABLE_ROLES), None)


def _parse_enum(enum_cls: Any, val: Any) -> Any | None:
    s = str(val or "").strip().lower().replace(" ", "_")
    return next((e for e in enum_cls if e.value == s), None)


def _number(val: Any) -> float | None:
    if val is None or str(val).strip() == "":
        return None
    try:
        return float(str(val).replace(",", ""))
    except ValueError:
        raise ValueError(f"Not a number: {val}")


async def _find_shift(session: AsyncSession, dept: Department | None, code: str) -> Shift | None:
    query = select(Shift).where(Shift.code == code.strip().upper())
    if dept is not None:
        query = query.where(Shift.plant_id == dept.plant_id)
    return (await session.execute(query)).scalars().first()


async def validate_row(session: AsyncSession, actor: User, row: dict[str, Any], *, org_id: UUID) -> list[str]:
    errors: list[str] = []
    for col in REQUIRED:
        if row.get(col) in (None, ""):
            errors.append(f"Missing required field: {col}")

    role = _parse_role(row.get("role"))
    if row.get("role") and not role:
        errors.append(f"Invalid role: {row.get('role')} (use one of {', '.join(IMPORTABLE_ROLES)})")

    dept, dept_errors = await _resolve_department(session, org_id, row.get("department_id"))
    errors.extend(dept_errors)
    if role and role.value not in NO_DEPARTMENT_ROLES and not row.get("department_id"):
        errors.append("Missing required field: department (needed for everyone except HR)")
    if role and role.value == "maintenance" and not _parse_enum(ObservationCategory, row.get("maintenance_division")):
        errors.append("maintenance_division is required for maintenance staff (quality, safety, energy, equipment or process)")
    if row.get("maintenance_division") and not _parse_enum(ObservationCategory, row.get("maintenance_division")):
        errors.append(f"Invalid maintenance_division: {row.get('maintenance_division')}")

    email = str(row.get("email", "")).strip().lower()
    if email:
        if not EMAIL_RE.match(email):
            errors.append(f"Invalid email: {email}")
        elif (await session.execute(select(User).where(User.email == email))).scalar_one_or_none():
            errors.append("Email already registered")

    if row.get("password") not in (None, "") and len(str(row["password"])) < 6:
        errors.append("Password must be at least 6 characters")

    uid = str(row.get("employee_uid") or "").strip()
    if uid and (await session.execute(select(User).where(User.employee_uid == uid))).scalar_one_or_none():
        errors.append("employee_uid already exists")

    if row.get("date_of_joining") and _parse_date(row.get("date_of_joining")) is None:
        errors.append("Invalid date_of_joining (use YYYY-MM-DD or DD-MM-YYYY)")
    if row.get("employment_type") and not _parse_enum(EmploymentType, row.get("employment_type")):
        errors.append("Invalid employment_type (permanent, contract or temporary)")
    if row.get("employment_status") and not _parse_enum(EmploymentStatus, row.get("employment_status")):
        errors.append("Invalid employment_status (active, on_leave, resigned or terminated)")

    if row.get("shift"):
        if not dept:
            errors.append("A shift needs a department")
        elif not await _find_shift(session, dept, str(row["shift"])):
            errors.append(f"Shift not found: {row['shift']} (A, B or C)")

    for col in ("basic", "hra", "allowances", "pf", "esi"):
        try:
            _number(row.get(col))
        except ValueError as exc:
            errors.append(f"{col}: {exc}")
    if any(row.get(c) not in (None, "") for c in ("hra", "allowances", "pf", "esi")) and row.get("basic") in (None, ""):
        errors.append("basic is required when any salary column is filled")

    return errors


async def import_row(session: AsyncSession, actor: User, row: dict[str, Any], *, org_id: UUID) -> dict[str, Any]:
    role = _parse_role(row["role"])
    if not role:
        raise ValueError("Invalid role")
    dept, _ = await _resolve_department(session, org_id, row.get("department_id"))
    emp_type = _parse_enum(EmploymentType, row.get("employment_type"))
    joined = _parse_date(row.get("date_of_joining"))

    data = WorkforceEmployeeCreate(
        email=str(row["email"]).strip().lower(),
        password=str(row["password"]),
        full_name=str(row["full_name"]).strip(),
        role=role,
        department_id=dept.id if dept else None,
        plant_id=dept.plant_id if dept else None,
        designation=str(row["designation"]).strip() if row.get("designation") else None,
        phone=str(row["phone"]).strip() if row.get("phone") else None,
        maintenance_division=_parse_enum(ObservationCategory, row.get("maintenance_division")),
        employment_status=_parse_enum(EmploymentStatus, row.get("employment_status")) or EmploymentStatus.ACTIVE,
        date_of_joining=joined,
        employment_type=emp_type,
    )
    profile = await _workforce.create_employee(session, actor, data)
    user = await session.get(User, profile.id)
    if user and row.get("employee_uid"):
        user.employee_uid = str(row["employee_uid"]).strip()
        await session.flush()

    if row.get("shift") and dept:
        shift = await _find_shift(session, dept, str(row["shift"]))
        if shift:
            await _workforce.create_shift_assignment(
                session, actor,
                ShiftAssignmentCreate(user_id=profile.id, department_id=dept.id, shift_id=shift.id, effective_date=joined or date.today()),
            )
    basic = _number(row.get("basic"))
    if basic:
        await _payroll.create_salary_structure(
            session, actor,
            SalaryStructureCreate(
                user_id=profile.id, basic=basic, hra=_number(row.get("hra")) or 0,
                allowances=_number(row.get("allowances")) or 0, pf=_number(row.get("pf")) or 0,
                esi=_number(row.get("esi")) or 0, effective_from=joined or date.today(),
            ),
        )
    return {"entity_id": profile.id, "action": "created"}


# --------------------------------------------------------------------------- template
HEADER_NOTES = {
    "employee_uid": "Your own employee code (optional, must be unique)",
    "full_name": "Required",
    "email": "Required. Used to log in. Must be unique",
    "password": "Required, at least 6 characters. Tell staff to keep it private",
    "role": "Required. Pick from the list",
    "department_id": "Department code from the Lists sheet. Needed for everyone except HR",
    "maintenance_division": "Only for maintenance staff: quality, safety, energy, equipment or process",
    "shift": "A, B or C (optional). Creates the shift assignment",
    "designation": "Job title (optional)",
    "phone": "Optional",
    "employment_type": "permanent, contract or temporary (optional)",
    "employment_status": "active (default), on_leave, resigned or terminated",
    "date_of_joining": "YYYY-MM-DD or DD-MM-YYYY (optional)",
    "basic": "Monthly basic pay (optional; fill to create the salary structure)",
    "hra": "Monthly HRA (optional)",
    "allowances": "Monthly allowances (optional)",
    "pf": "Monthly PF deduction (optional)",
    "esi": "Monthly ESI deduction (optional)",
}

EXAMPLES = [
    ["CS-1001", "Ramesh Kumar", "ramesh.kumar@example.com", "Temp@1234", "worker", "SMS", "", "A", "Furnace operator", "9876543210", "permanent", "active", "2024-04-01", 22000, 7000, 1500, 1800, 500],
    ["CS-1002", "Sunita Devi", "sunita.devi@example.com", "Temp@1234", "supervisor", "ROLL", "", "B", "Shift in-charge", "9876501234", "permanent", "active", "2022-07-15", 32000, 10000, 3000, 2400, 600],
    ["CS-1003", "Anil Sharma", "anil.sharma@example.com", "Temp@1234", "maintenance", "SMS", "equipment", "", "Fitter", "9876512345", "contract", "active", "2025-01-10", "", "", "", "", ""],
]


async def build_template(session: AsyncSession, org_id: UUID) -> bytes:
    depts = (await session.execute(select(Department).where(Department.organisation_id == org_id).order_by(Department.code))).scalars().all()
    wb = Workbook()

    guide = wb.active
    guide.title = "Read me first"
    lines = [
        ("Employee upload template", True),
        ("", False),
        ("1. Fill one row per person on the 'Employees' sheet. Do not rename or reorder the column headings.", False),
        ("2. Required columns have a dark heading: full_name, email, password, role. Everything else is optional.", False),
        ("3. department: use the short code from the 'Lists' sheet (for example SMS). Needed for everyone except HR staff.", False),
        ("4. The 'Example' sheet shows three filled rows. It is only an illustration and is never read on upload.", False),
        ("5. Save as .xlsx and upload it in the app: Workforce > Employees > Upload Excel.", False),
        ("6. The app checks every row first and lists problems per row. Nothing is created until you confirm.", False),
        ("7. Rows with problems are skipped; you can fix them and upload just those rows again.", False),
        ("8. Passwords are stored securely once uploaded. Delete this file afterwards because it contains passwords.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = guide.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=14 if bold else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 110

    ws = wb.create_sheet("Employees")
    required_fill = PatternFill("solid", fgColor="1F3A5F")
    optional_fill = PatternFill("solid", fgColor="6B7C93")
    for idx, col in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = required_fill if col in REQUIRED else optional_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = max(16, len(col) + 4)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    for col in ("employee_uid", "phone", "password"):
        for r in range(2, 1002):
            ws.cell(row=r, column=COLUMNS.index(col) + 1).number_format = "@"  # keep leading zeros and long numbers as typed

    def dropdown(column: str, formula: str) -> None:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        letter = get_column_letter(COLUMNS.index(column) + 1)
        dv.add(f"{letter}2:{letter}1000")

    lists = wb.create_sheet("Lists")
    lists.append(["role", "department_code", "department_name", "shift", "employment_type", "employment_status", "maintenance_division"])
    for c in lists[1]:
        c.font = Font(bold=True)
    statuses = [e.value for e in EmploymentStatus]
    types = [e.value for e in EmploymentType]
    divisions = [e.value for e in ObservationCategory]
    for i in range(max(len(IMPORTABLE_ROLES), len(depts), 3, len(statuses), len(divisions))):
        lists.append([
            IMPORTABLE_ROLES[i] if i < len(IMPORTABLE_ROLES) else None,
            depts[i].code if i < len(depts) else None,
            depts[i].name if i < len(depts) else None,
            ["A", "B", "C"][i] if i < 3 else None,
            types[i] if i < len(types) else None,
            statuses[i] if i < len(statuses) else None,
            divisions[i] if i < len(divisions) else None,
        ])
    for letter in "ABCDEFG":
        lists.column_dimensions[letter].width = 22
    dropdown("role", f"=Lists!$A$2:$A${1 + len(IMPORTABLE_ROLES)}")
    if depts:
        dropdown("department_id", f"=Lists!$B$2:$B${1 + len(depts)}")
    dropdown("shift", "=Lists!$D$2:$D$4")
    dropdown("employment_type", f"=Lists!$E$2:$E${1 + len(types)}")
    dropdown("employment_status", f"=Lists!$F$2:$F${1 + len(statuses)}")
    dropdown("maintenance_division", f"=Lists!$G$2:$G${1 + len(divisions)}")

    ex = wb.create_sheet("Example")
    ex.append(COLUMNS)
    for c in ex[1]:
        c.font = Font(bold=True)
    for row in EXAMPLES:
        ex.append(row)
    ex.append([])
    ex.append(["Notes"] + [HEADER_NOTES[c] for c in COLUMNS[1:]])
    for idx in range(1, len(COLUMNS) + 1):
        ex.column_dimensions[get_column_letter(idx)].width = 24
    ex.cell(row=ex.max_row, column=1).font = Font(bold=True)
    for c in ex[ex.max_row]:
        c.alignment = Alignment(wrap_text=True, vertical="top")

    wb.move_sheet("Employees", offset=-1)  # open on the sheet people fill in
    wb.active = wb.sheetnames.index("Employees")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def register() -> None:
    register_module(
        "employees",
        columns=COLUMNS,
        required=REQUIRED,
        validate_row=validate_row,
        import_row=import_row,
        aliases=ALIASES,
        unique_keys=["email", "employee_uid"],
        build_template=build_template,
    )


register()
